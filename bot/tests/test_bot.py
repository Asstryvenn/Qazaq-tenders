"""Тесты представления и вспомогательных модулей бота (без сети и без Telegram)."""
import asyncio
import json
from io import BytesIO
from pathlib import Path

import pytest

from calculator import ALL_LEVERS_MASK, analyze_tender, scenario_from_mask
from cards import full_breakdown, lot_card, lot_keyboard, search_keyboard, search_results, simulator_keyboard, simulator_view, twin_view
from documents import DocumentError, _assert_public_host, _chunk, extract_docx, find_url
from models import CompanyTwin, SearchQuery, TenderSpec
from storage import Storage
from tenders import filter_lots, find_by_url, heuristic_search, lot_key, match_city
from texts import money, pct

LOTS = [TenderSpec.model_validate(x) for x in json.loads((Path(__file__).resolve().parents[1] / "data" / "sample_tenders.json").read_text("utf-8"))["tenders"]]
TWIN = CompanyTwin.from_answers(25, "astana", 10, "construction", "simplified", name="ТОО «Тест»")


def _callback_sizes(markup):
    for row in markup.inline_keyboard:
        for b in row:
            if b.callback_data:
                assert len(b.callback_data.encode("utf-8")) <= 64, b.callback_data


@pytest.mark.parametrize("lang", ["kz", "ru"])
@pytest.mark.parametrize("spec", LOTS, ids=lambda s: s.id)
def test_views_render_within_telegram_limits(spec, lang):
    key = lot_key(spec.id)
    base = analyze_tender(spec, TWIN)
    sim = analyze_tender(spec, TWIN, scenario_from_mask(ALL_LEVERS_MASK))
    for text in (
        lot_card(spec, base, TWIN, lang),
        full_breakdown(spec, base, TWIN, lang),
        simulator_view(spec, base, sim, ALL_LEVERS_MASK, TWIN, lang),
        simulator_view(spec, base, base, 0, TWIN, lang),
    ):
        assert 0 < len(text) < 4096
        assert "{" not in text  # все плейсхолдеры подставлены
    _callback_sizes(lot_keyboard(key, spec, lang))
    _callback_sizes(simulator_keyboard(key, ALL_LEVERS_MASK, lang))


def test_card_contains_required_fields():
    spec = LOTS[0]
    text = lot_card(spec, analyze_tender(spec, TWIN), TWIN, "ru")
    for needle in ("Заказчик", "TOS", "Чистая прибыль", "Кассовый разрыв", "Логистика", "Главный риск"):
        assert needle in text
    kz = lot_card(spec, analyze_tender(spec, TWIN), TWIN, "kz")
    assert "Тапсырыс беруші" in kz


def test_search_views():
    rows = [(lot_key(s.id), s, analyze_tender(s, TWIN)) for s in LOTS[:5]]
    q = SearchQuery(keywords=["ремонт"], city_id="astana", max_amount_kzt=10e6)
    assert "Найдено" in search_results(rows, q, TWIN, "ru", live=False)
    _callback_sizes(search_keyboard(rows, "kz"))
    assert "Цифровой двойник" in twin_view(TWIN, "ru")


def test_heuristic_search_ru_voice_example():
    q = heuristic_search("Найди мне тендеры на мебель в Астане до 10 млн тенге")
    assert q.city_id == "astana" and q.max_amount_kzt == 10_000_000 and q.industry == "furniture"


def test_heuristic_search_kz():
    q = heuristic_search("Алматыда 50 млн теңгеге дейін компьютер жеткізу")
    assert q.city_id == "almaty" and q.max_amount_kzt == 50_000_000 and q.industry == "it"
    q2 = heuristic_search("ремонт от 20 млн")
    assert q2.min_amount_kzt == 20_000_000 and q2.max_amount_kzt is None


def test_filter_lots():
    q = SearchQuery(industry="construction")
    titles = [s.title for s in filter_lots(LOTS, q)]
    assert any("кровли" in t for t in titles) and any("арматуры" in t for t in titles)
    assert all(s.contract_amount <= 60e6 for s in filter_lots(LOTS, SearchQuery(max_amount_kzt=60e6)))


def test_match_city_forms():
    assert match_city("в Караганде") == "karaganda"
    assert match_city("Қарағандыда") == "karaganda"
    assert match_city("Астанада") == "astana"
    assert match_city("хорошая погода") is None


def test_find_by_url_and_find_url():
    spec = LOTS[0]
    assert find_by_url(LOTS, f"https://goszakup.gov.kz/ru/announce/index/{spec.external_id}") is spec
    assert find_url("смотри https://goszakup.gov.kz/ru/lot/123, спасибо") == "https://goszakup.gov.kz/ru/lot/123"


def test_money_format():
    assert money(84_000_000, "ru") == "84 млн ₸"
    assert money(12_400_000, "ru") == "12,4 млн ₸"
    assert money(520_000, "kz") == "520 мың ₸"
    assert money(-2_100_000, "ru") == "−2,1 млн ₸"
    assert pct(6.25) == "6,2%" or pct(6.25) == "6,3%"


def test_docx_extraction():
    import docx

    d = docx.Document()
    d.add_paragraph("Техническая спецификация. Сумма закупки 10 000 000 тенге.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Оплата"
    table.rows[0].cells[1].text = "в течение 45 дней"
    buf = BytesIO()
    d.save(buf)
    pages = extract_docx(buf.getvalue())
    text = " ".join(p.text for p in pages)
    assert "10 000 000" in text and "45 дней" in text


def test_chunking_pages():
    pages = _chunk("\n".join(["абзац " * 100] * 20), size=1000)
    assert len(pages) > 5 and all(len(p.text) <= 1000 for p in pages)
    assert [p.page for p in pages] == list(range(1, len(pages) + 1))


def test_ssrf_guard_blocks_private_hosts():
    for host in ("localhost", "127.0.0.1", "10.0.0.5", "169.254.169.254"):
        with pytest.raises(DocumentError):
            asyncio.run(_assert_public_host(host))


def test_storage_roundtrip_and_ai_limit(tmp_path):
    async def scenario():
        st = Storage(str(tmp_path / "t.db"))
        await st.ensure_user(1, "kz")
        await st.save_twin(1, TWIN)
        user = await st.get_user(1)
        assert user["twin"] == TWIN and user["lang"] == "kz"
        key = lot_key(LOTS[0].id)
        await st.put_lot(key, LOTS[0])
        assert (await st.get_lot(key)) == LOTS[0]
        await st.hide(1, key)
        assert key in await st.hidden_keys(1)
        assert [await st.use_ai(1, 2) for _ in range(3)] == [True, True, False]

    asyncio.run(scenario())
