"""
Извлечение текста из PDF / DOCX / веб-страницы лота, постранично.

Номера страниц важны: AI указывает, на какой странице найден рискованный пункт.
У DOCX и HTML нет «страниц», поэтому текст режется на условные страницы по ~3000 символов.

Загрузка ссылок защищена от SSRF: разрешены только http/https и только публичные IP-адреса.
Каждый редирект проверяется заново.
"""
from __future__ import annotations

import asyncio
import ipaddress
import re
import socket
from html.parser import HTMLParser
from io import BytesIO
from typing import List, Tuple
from urllib.parse import urljoin, urlparse

import httpx

from models import SpecPage

MAX_FILE_BYTES = 20 * 1024 * 1024  # потолок getFile в Telegram Bot API
MAX_URL_BYTES = 10 * 1024 * 1024
PSEUDO_PAGE_CHARS = 3000
MIN_TEXT_CHARS = 200  # меньше — вероятно, скан без текстового слоя
USER_AGENT = "QazaqTendersBot/1.0 (+https://www.qazaqtenders.kz)"


class DocumentError(Exception):
    """code: unsupported | too_big | no_text | bad_url | blocked_host | http | redirects | broken."""

    def __init__(self, code: str, detail: str = "") -> None:
        super().__init__(f"{code}: {detail}" if detail else code)
        self.code = code
        self.detail = detail


# ------------------------------- разбор файлов ------------------------------- #


def _chunk(text: str, size: int = PSEUDO_PAGE_CHARS) -> List[SpecPage]:
    """Режет сплошной текст на условные страницы, стараясь резать по абзацам."""
    pages: List[SpecPage] = []
    buf = ""
    for para in re.split(r"\n{1,}", text):
        para = para.strip()
        if not para:
            continue
        if buf and len(buf) + len(para) + 1 > size:
            pages.append(SpecPage(page=len(pages) + 1, text=buf))
            buf = ""
        buf = f"{buf}\n{para}" if buf else para
        while len(buf) > size:  # очень длинный абзац
            pages.append(SpecPage(page=len(pages) + 1, text=buf[:size]))
            buf = buf[size:]
    if buf:
        pages.append(SpecPage(page=len(pages) + 1, text=buf))
    return pages


def extract_pdf(data: bytes) -> List[SpecPage]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")  # многие PDF «зашифрованы» пустым паролем
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            text = re.sub(r"[ \t]+", " ", page.extract_text() or "").strip()
            pages.append(SpecPage(page=i, text=text))
        return pages
    except (PdfReadError, ValueError, KeyError) as e:
        raise DocumentError("broken", str(e)) from e


def extract_docx(data: bytes) -> List[SpecPage]:
    import docx  # python-docx

    try:
        document = docx.Document(BytesIO(data))
    except Exception as e:  # noqa: BLE001 — python-docx бросает разные исключения
        raise DocumentError("broken", str(e)) from e
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:  # условия оплаты часто спрятаны в таблицах
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _chunk("\n".join(parts))


class _TextFromHTML(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "table", "section"}

    def __init__(self) -> None:
        super().__init__()
        self.out: List[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data):  # noqa: ANN001
        if not self._skip:
            self.out.append(data)


def extract_html(html_text: str) -> List[SpecPage]:
    parser = _TextFromHTML()
    parser.feed(html_text)
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(parser.out))
    text = re.sub(r"\n\s*\n+", "\n", text)
    return _chunk(text)


def detect_kind(file_name: str, mime: str = "", head: bytes = b"") -> str:
    """pdf | docx | html | unknown — по сигнатуре, MIME-типу и расширению."""
    name = (file_name or "").lower()
    mime = (mime or "").lower()
    if head.startswith(b"%PDF") or "pdf" in mime or name.endswith(".pdf"):
        return "pdf"
    if "wordprocessingml" in mime or name.endswith(".docx"):
        return "docx"
    if "html" in mime or name.endswith((".html", ".htm")):
        return "html"
    return "unknown"


def extract_pages(data: bytes, file_name: str = "", mime: str = "") -> List[SpecPage]:
    """Любой поддерживаемый файл → страницы. Бросает DocumentError."""
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError("too_big")
    kind = detect_kind(file_name, mime, data[:8])
    if kind == "pdf":
        pages = extract_pdf(data)
    elif kind == "docx":
        pages = extract_docx(data)
    elif kind == "html":
        pages = extract_html(data.decode("utf-8", errors="replace"))
    else:
        raise DocumentError("unsupported", file_name or mime)
    if total_chars(pages) < MIN_TEXT_CHARS:
        raise DocumentError("no_text")
    return pages


def total_chars(pages: List[SpecPage]) -> int:
    return sum(len(p.text) for p in pages)


# ------------------------- безопасная загрузка ссылок ------------------------ #

URL_RE = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)


def find_url(text: str) -> str:
    m = URL_RE.search(text or "")
    return m.group(0).rstrip(").,;»") if m else ""


async def _assert_public_host(host: str) -> None:
    """Запрещает внутренние адреса (localhost, 10.x, 192.168.x, 169.254.x, …), чтобы
    бот нельзя было использовать для запросов во внутреннюю сеть сервера."""
    loop = asyncio.get_running_loop()
    try:
        infos = await loop.getaddrinfo(host, None, type=socket.SOCK_STREAM)
    except socket.gaierror as e:
        raise DocumentError("bad_url", host) from e
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global:
            raise DocumentError("blocked_host", host)


async def fetch_url(url: str) -> Tuple[bytes, str, str]:
    """Скачивает страницу или файл по ссылке. Возвращает (данные, content-type, итоговый URL)."""
    async with httpx.AsyncClient(timeout=20, follow_redirects=False, headers={"User-Agent": USER_AGENT}) as client:
        for _ in range(5):
            parsed = urlparse(url)
            if parsed.scheme not in ("http", "https") or not parsed.hostname:
                raise DocumentError("bad_url", url)
            await _assert_public_host(parsed.hostname)
            async with client.stream("GET", url) as r:
                if r.status_code in (301, 302, 303, 307, 308):
                    url = urljoin(url, r.headers.get("location", ""))
                    continue
                if r.status_code >= 400:
                    raise DocumentError("http", str(r.status_code))
                chunks: List[bytes] = []
                size = 0
                async for chunk in r.aiter_bytes():
                    size += len(chunk)
                    if size > MAX_URL_BYTES:
                        raise DocumentError("too_big")
                    chunks.append(chunk)
                return b"".join(chunks), r.headers.get("content-type", ""), str(r.url)
    raise DocumentError("redirects")


async def pages_from_url(url: str) -> List[SpecPage]:
    data, ctype, final_url = await fetch_url(url)
    name = urlparse(final_url).path.rsplit("/", 1)[-1]
    kind = detect_kind(name, ctype, data[:8])
    return extract_pages(data, name if kind != "unknown" else "page.html", ctype if kind != "unknown" else "text/html")
