from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "output" / "pdf"
WORK = Path(__file__).resolve().parent
OUT_DOCX = OUT_DIR / "Qazaq_Tenders_gylimi_zhoba_Nurali_Baibosynov.docx"

sys.path.insert(0, str(ROOT / "bot"))
from calculator import analyze_tender  # noqa: E402
from models import CompanyTwin, Scenario, TenderSpec  # noqa: E402


NAVY = "17365D"
BLUE = "1F4E78"
PALE = "EAF2F8"
LIGHT = "F5F7FA"
GRAY = "666666"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep(paragraph, next_: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)
    if next_:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def set_font(run, name="Times New Roman", size=12, bold=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    set_font(run, size=10, color=RGBColor(90, 90, 90))


def add_body(doc: Document, text: str, *, bold_lead: str | None = None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.0) if align == WD_ALIGN_PARAGRAPH.JUSTIFY else Cm(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead) :])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc: Document, items: list[str], level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Cm(0.7 + level * 0.6)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        set_font(p.add_run(item), size=11.5)
        set_keep(p)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        set_font(p.add_run(item), size=11.5)
        set_keep(p)


def add_heading(doc: Document, text: str, level=1, new_page=False) -> None:
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=15 if level == 1 else 13, bold=True)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=10.5, bold=True)


def add_picture(doc: Document, path: Path, width_cm: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _m(tag: str, text: str | None = None):
    node = OxmlElement(f"m:{tag}")
    if text is not None:
        t = OxmlElement("m:t")
        t.text = text
        node.append(t)
    return node


def _mr(text: str):
    return _m("r", text)


def _sub(base: str, subscript: str):
    node = _m("sSub")
    e = _m("e")
    e.append(_mr(base))
    s = _m("sub")
    s.append(_mr(subscript))
    node.extend([e, s])
    return node


def _frac(numerator: list, denominator: list):
    node = _m("f")
    n = _m("num")
    d = _m("den")
    for part in numerator:
        n.append(part)
    for part in denominator:
        d.append(part)
    node.extend([n, d])
    return node


def add_equation(doc: Document, parts: list, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    math_para = _m("oMathPara")
    math = _m("oMath")
    for part in parts:
        math.append(part)
    math_para.append(math)
    p._p.append(math_para)
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    table.rows[0].height = None
    set_repeat_table_header(table.rows[0])
    for i, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cell)
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value, width) in enumerate(zip(cells, row, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            set_cell_border(cell)
            if ridx % 2:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(value) > 24 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(value), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def make_visuals() -> tuple[dict, dict[str, object]]:
    data = json.loads((ROOT / "bot" / "data" / "sample_tenders.json").read_text("utf-8"))
    spec = TenderSpec.model_validate(data["tenders"][0])
    company = CompanyTwin.from_answers(52, "astana", 10, "it", "general", name="Үлгі кәсіпорын")
    scenarios = {
        "Автоматты таңдау": Scenario(),
        "Автофура": Scenario(transport_mode="truck"),
        "Теміржол": Scenario(transport_mode="rail"),
        "Авиа": Scenario(transport_mode="air"),
        "30% аванс": Scenario(advance_percentage_override=30, verified_fields=["advance_percentage"]),
    }
    results = {name: analyze_tender(spec, company, scenario) for name, scenario in scenarios.items()}

    regular = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(regular, 31)
    font_bold = ImageFont.truetype(bold, 31)
    font_small = ImageFont.truetype(regular, 25)

    img = Image.new("RGB", (1900, 650), "white")
    draw = ImageDraw.Draw(img)
    boxes = [
        (50, 170, 390, 420, "Дереккөздер\nPDF  DOCX  API", "#EAF2F8"),
        (520, 170, 870, 420, "AI қабаты\nфактілер және JSON", "#D6EAF8"),
        (1000, 170, 1430, 420, "Экономикалық қабат\nпайда  CF  TOS", "#D5F5E3"),
        (1550, 170, 1850, 420, "Нәтиже\nвеб және Telegram", "#FDEBD0"),
    ]
    for x1, y1, x2, y2, label, fill in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=fill, outline="#1F4E78", width=5)
        bbox = draw.multiline_textbbox((0, 0), label, font=font_bold, spacing=18, align="center")
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), label, font=font_bold, fill="#17202A", spacing=18, align="center")
    for x1, x2 in ((390, 520), (870, 1000), (1430, 1550)):
        draw.line((x1 + 12, 295, x2 - 26, 295), fill="#1F4E78", width=7)
        draw.polygon(((x2 - 26, 280), (x2 - 26, 310), (x2 - 2, 295)), fill="#1F4E78")
    draw.text((570, 505), "ЖИ есеп шығармайды", font=font_small, fill="#7B241C")
    draw.text((1035, 505), "Формулалар детерминирленген", font=font_small, fill="#196F3D")
    img.save(WORK / "architecture.png")

    baseline = results["Автоматты таңдау"]
    advance = results["30% аванс"]
    chart = Image.new("RGB", (1800, 850), "white")
    d = ImageDraw.Draw(chart)
    left, top, right, bottom = 170, 80, 1730, 700
    values = [p.balance / 1e6 for p in baseline.timeline] + [p.balance / 1e6 for p in advance.timeline]
    lo, hi = min(values) - 2, max(values) + 2
    def xy(day, value, n):
        return left + day / max(1, n - 1) * (right - left), bottom - (value - lo) / (hi - lo) * (bottom - top)
    for j in range(6):
        value = lo + (hi - lo) * j / 5
        y = xy(0, value, 2)[1]
        d.line((left, y, right, y), fill="#D9E1E8", width=2)
        d.text((25, y - 14), f"{value:.0f}", font=font_small, fill="#566573")
    for series, color in ((baseline.timeline, "#C0392B"), (advance.timeline, "#1E8449")):
        pts = [xy(p.day, p.balance / 1e6, len(series)) for p in series]
        d.line(pts, fill=color, width=7, joint="curve")
    zero_y = xy(0, 0, 2)[1]
    d.line((left, zero_y, right, zero_y), fill="#2C3E50", width=3)
    d.line((left, top, left, bottom), fill="#2C3E50", width=3)
    d.line((left, bottom, right, bottom), fill="#2C3E50", width=3)
    d.text((20, 20), "Ақша қалдығы млн ₸", font=font_small, fill="#2C3E50")
    d.text((900, 750), "Күн", font=font_small, fill="#2C3E50")
    d.line((1200, 35, 1260, 35), fill="#C0392B", width=7)
    d.text((1280, 18), "Аванссыз", font=font_small, fill="#2C3E50")
    d.line((1470, 35, 1530, 35), fill="#1E8449", width=7)
    d.text((1550, 18), "30% аванс", font=font_small, fill="#2C3E50")
    chart.save(WORK / "cashflow_case.png")

    transport_names = ["Автофура", "Теміржол", "Авиа"]
    costs = [results[name].costs.logistics / 1e6 for name in transport_names]
    bar_img = Image.new("RGB", (1700, 720), "white")
    d = ImageDraw.Draw(bar_img)
    left, top, right, bottom = 160, 70, 1620, 590
    d.line((left, top, left, bottom), fill="#2C3E50", width=3)
    d.line((left, bottom, right, bottom), fill="#2C3E50", width=3)
    maxv = max(costs) * 1.15
    colors = ["#2874A6", "#239B56", "#AF7AC5"]
    for j in range(5):
        value = maxv * j / 4
        y = bottom - value / maxv * (bottom - top)
        d.line((left, y, right, y), fill="#D9E1E8", width=2)
        d.text((45, y - 14), f"{value:.0f}", font=font_small, fill="#566573")
    for i, (name, value, color) in enumerate(zip(transport_names, costs, colors)):
        x1 = 300 + i * 450
        x2 = x1 + 230
        y = bottom - value / maxv * (bottom - top)
        d.rectangle((x1, y, x2, bottom), fill=color)
        label = f"{value:.2f}"
        tw = d.textbbox((0, 0), label, font=font_bold)[2]
        d.text(((x1 + x2 - tw) / 2, y - 48), label, font=font_bold, fill="#17202A")
        tw = d.textbbox((0, 0), name, font=font_small)[2]
        d.text(((x1 + x2 - tw) / 2, bottom + 20), name, font=font_small, fill="#17202A")
    d.text((20, 18), "Логистика құны млн ₸", font=font_small, fill="#2C3E50")
    bar_img.save(WORK / "logistics_case.png")

    return results, {"spec": spec, "company": company}


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.8)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for name, size in (("Title", 16), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("Qazaq Tenders ғылыми жобасы"), size=9, color=RGBColor(105, 105, 105))
    page_field(section.footer.paragraphs[0])


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("ҚАЗАҚСТАН РЕСПУБЛИКАСЫНЫҢ ОҚУ АҒАРТУ МИНИСТРЛІГІ"), size=11, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("АСТАНА ҚАЛАСЫ НҰРА АУДАНЫНДАҒЫ\nЖАРАТЫЛЫСТАНУ МАТЕМАТИКА БАҒЫТЫНДАҒЫ\nНАЗАРБАЕВ ЗИЯТКЕРЛІК МЕКТЕБІ"), size=11, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_font(p.add_run("ҒЫЛЫМИ ЖОБА"), size=17, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    set_font(
        p.add_run(
            "ҚАЗАҚСТАНДАҒЫ ШАҒЫН ЖӘНЕ ОРТА БИЗНЕСТІҢ\n"
            "МЕМЛЕКЕТТІК САТЫП АЛУҒА ҚАТЫСУЫНЫҢ\n"
            "ЭКОНОМИКАЛЫҚ ТИІМДІЛІГІН БАҒАЛАУҒА\n"
            "АРНАЛҒАН ИНТЕЛЛЕКТУАЛДЫ МОДЕЛЬ"
        ),
        size=15,
        bold=True,
    )
    for _ in range(4):
        doc.add_paragraph()
    rows = [
        ("Орындаған", "Байбосынов Нұрали Айбарұлы"),
        ("Ғылыми жетекші", "Нургалиева Гүлжан Әбенқызы"),
    ]
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    table.autofit = False
    for r, row in zip(table.rows, rows):
        r.cells[0].width = Cm(4)
        r.cells[1].width = Cm(7)
        for c in r.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{edge}")
                e.set(qn("w:val"), "nil")
                borders.append(e)
            tc_pr.append(borders)
        set_font(r.cells[0].paragraphs[0].add_run(row[0] + ":"), size=11, bold=True)
        set_font(r.cells[1].paragraphs[0].add_run(row[1]), size=11)
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Астана 2026"), size=12)


def add_front_matter(doc: Document) -> None:
    add_heading(doc, "АҢДАТПА", new_page=True)
    add_body(doc, "Бұл зерттеуде Қазақстандағы шағын және орта бизнес субъектілерінің мемлекеттік сатып алуға қатысуы нақты кәсіпорынның қаржылық мүмкіндіктеріне байланысты бағаланды. Зерттеу нәтижесі ретінде Qazaq Tenders бағдарламалық прототипі жасалды. Жүйе техникалық ерекшеліктен шарттарды құрылымдайды, кәсіпорынның цифрлық бейінімен салыстырады және пайда, маржа, күнделікті ақша ағыны, логистика мен құқықтық тәуекелдер негізінде Tender Opportunity Score индексін есептейді.")
    add_body(doc, "Жұмыстың негізгі әдістемелік қағидасы жасанды интеллект пен қаржылық есептеуді бөлуге негізделген. Тілдік модель PDF және DOCX құжаттарынан фактілерді ғана шығарады. Ақшалай көрсеткіштер қайталанатын нәтижелер беретін детерминирленген алгоритмде есептеледі. Бағдарламалық валидация 210 салыстырмалы сценарийді және барлығы 450 автоматтандырылған тестті қамтыды. Бұл тесттер кодтың тұрақтылығын дәлелдейді, бірақ экономикалық болжамның 95 пайыздық дәлдігін білдірмейді.")
    add_body(doc, "Үлгілік кейс аванссыз шартта пайданың оң болғанымен, 13,2 млн теңге кассалық алшақтық туындайтынын көрсетті. 30 пайыз аванс енгізілген сценарийде алшақтық жойылып, TOS 62,7 балдан 73,4 балға өсті. Нәтиже тендердің тартымдылығы келісімшарт сомасымен ғана емес, төлем күнтізбесі мен кәсіпорын капиталының сәйкестігімен анықталатынын көрсетеді.")
    add_body(doc, "Түйін сөздер: мемлекеттік сатып алу, шағын және орта бизнес, ақша ағыны, кассалық алшақтық, экономикалық модель, TOS, NLP, логистика, сценарийлік талдау.", bold_lead="Түйін сөздер:", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "МАЗМҰНЫ", new_page=True)
    toc_pages = json.loads(os.getenv("TOC_PAGES_JSON", "{}"))
    toc = [
        "КІРІСПЕ",
        "1 МЕМЛЕКЕТТІК САТЫП АЛУДАҒЫ ШОБ МӘСЕЛЕСІ",
        "2 ЗЕРТТЕУ ӘДІСТЕМЕСІ",
        "3 QAZAQ TENDERS ЖҮЙЕСІНІҢ АРХИТЕКТУРАСЫ",
        "4 ЭКОНОМИКАЛЫҚ МАТЕМАТИКАЛЫҚ МОДЕЛЬ",
        "5 ЛОГИСТИКА ЖӘНЕ ЖЕТКІЗУШІЛЕР МОДУЛІ",
        "6 ЕСЕПТІК КЕЙС ЖӘНЕ СЦЕНАРИЙЛІК ТАЛДАУ",
        "7 БАҒДАРЛАМАЛЫҚ ВАЛИДАЦИЯ",
        "8 ШЕКТЕУЛЕР ЖӘНЕ ДАМЫТУ БАҒЫТТАРЫ",
        "ҚОРЫТЫНДЫ",
        "ПАЙДАЛАНЫЛҒАН ӘДЕБИЕТТЕР",
        "ҚОСЫМША",
    ]
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for title in toc:
        cells = table.add_row().cells
        cells[0].width = Cm(14.7)
        cells[1].width = Cm(1.2)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run(title), size=11)
        p = cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(p.add_run(str(toc_pages.get(title, "—"))), size=11)


def add_introduction(doc: Document) -> None:
    add_heading(doc, "КІРІСПЕ", new_page=True)
    add_heading(doc, "Зерттеудің өзектілігі", level=2)
    add_body(doc, "Мемлекеттік сатып алу Қазақстан экономикасындағы ірі институционалдық нарықтардың бірі. Қазақстан Республикасы Қаржы министрлігінің дерегі бойынша 2025 жылы мемлекеттік сатып алу көлемі 8,1 трлн теңгеге жетті [2]. Ұлттық статистика бюросы 2025 жылдың қорытындысында елде 2 175,6 мың белсенді шағын және орта кәсіпкерлік субъектісі болғанын, бұл секторда 4 527,4 мың адам жұмыс істегенін және ШОБ үлесі жалпы ішкі өнімнің 40,9 пайызын құрағанын хабарлады [3]. Сондықтан мемлекеттік тапсырысқа қауіпсіз қатысу ШОБ тұрақтылығына тікелей әсер етеді.")
    add_body(doc, "Тендердің жоғары сомасы оның нақты кәсіпорын үшін тиімді екенін өздігінен дәлелдемейді. Жеткізуші тауарды алдын ала сатып алуы, кепілдік қаражатын уақытша бұғаттауы, жеткізу мен салық шығындарын төлеуі мүмкін. Тапсырыс берушінің төлемі кейін түскен жағдайда бухгалтерлік пайда оң болғанымен, белгілі бір күні ақша қалдығы теріс мәнге түседі. Бұл жағдай кассалық алшақтық деп аталады.")
    add_body(doc, "Қолданыстағы цифрлық сервистер тендерді іздеу, техникалық ерекшелікті қысқарту немесе жеңіс бағасын болжау міндеттерін орындайды. Алайда ашық жарияланған сипаттамаларда кәсіпорынның нақты айналым капиталына негізделген күнделікті Cash Flow күнтізбесі әрқашан көрсетілмейді [15; 16]. Осы зерттеу дәл осы шешім қабылдау кезеңіне бағытталды.")
    add_heading(doc, "Зерттеу сұрағы мен мақсаты", level=2)
    add_body(doc, "Зерттеу сұрағы: тендер талаптары, кәсіпорынның айналым капиталы және орындау сценарийлері бір модельде біріктірілгенде, қатысу туралы шешімді қалай ашық әрі қайталанатын түрде есептеуге болады?")
    add_body(doc, "Зерттеудің мақсаты: Қазақстандағы ШОБ субъектілеріне арналған, техникалық ерекшеліктен алынған деректерді кәсіпорынның қаржылық бейінімен салыстырып, тендердің экономикалық орындылығын бағалайтын модель мен бағдарламалық прототип әзірлеу.")
    add_heading(doc, "Зерттеу міндеттері", level=2)
    add_numbered(doc, [
        "ШОБ үшін мемлекеттік сатып алуға қатысудың негізгі экономикалық тәуекелдерін анықтау.",
        "Нарықтағы сервистердің ашық жарияланған мүмкіндіктерін салыстыру.",
        "Техникалық ерекшеліктен фактілер мен жеткізушіге қойылатын талаптарды шығару тәсілін әзірлеу.",
        "Пайда, маржа, ақша ағыны, логистика және құқықтық тәуекелді бір есепте біріктіру.",
        "Сценарийлерді бір әрекетпен өзгертіп, нәтижені қайта есептейтін веб және Telegram интерфейсін құру.",
        "TypeScript және Python қозғалтқыштарының сәйкестігін автоматтандырылған тесттермен тексеру.",
        "Модельдің шектеулерін бөліп көрсету және эмпирикалық валидация жоспарын ұсыну.",
    ])
    add_heading(doc, "Гипотеза", level=2)
    add_body(doc, "Егер тендер тек келісімшарт сомасы бойынша емес, кәсіпорынның айналым капиталы, сатып алу құны, жеткізу күнтізбесі және тәуекелдері бойынша бағаланса, онда пайдасы оң болғанымен орындауға қауіпті лоттарды өтінім беруге дейін анықтауға болады.")
    add_heading(doc, "Ғылыми жаңалық пен практикалық маңыз", level=2)
    add_body(doc, "Жұмыстың жаңалығы жеке алгоритмнің әлемде алғаш рет ұсынылуы туралы талапқа емес, бірнеше әдістің бір шешім тізбегінде біріктірілуіне негізделеді. Qazaq Tenders құжаттық NLP талдауын, кәсіпорынның цифрлық бейінін, күнделікті Cash Flow симуляциясын және түсіндірілетін TOS индексін бір прототипте байланыстырады. Практикалық нәтиже ретінде жұмыс істейтін веб-қосымша мен Telegram-бот әзірленді.")


def add_market(doc: Document) -> None:
    add_heading(doc, "1 МЕМЛЕКЕТТІК САТЫП АЛУДАҒЫ ШОБ МӘСЕЛЕСІ", new_page=True)
    add_heading(doc, "1.1 Экономикалық тәуекелдер", level=2)
    add_body(doc, "ШОБ жеткізушісі үшін негізгі шектеу ақпараттың жоқтығынан гөрі оны экономикалық шешімге айналдыру қиындығында. Бір тендердің құжаттары бірнеше ондаған беттен тұруы мүмкін. Жеткізу мерзімі, төлем тәртібі, кепілдік, сертификат, тәжірибе және жабдық туралы талаптар әр бөлімде орналасады. Осы талаптардың біреуін өткізіп алу өтінімнің қабылданбауына немесе шартты орындай алмауға әкеледі.")
    add_bullets(doc, [
        "Қаржыландыру тәуекелі: сатып алу мен жеткізу төлемнен бұрын орындалады.",
        "Баға тәуекелі: болжамды өзіндік құн нақты коммерциялық ұсыныстан ауытқуы мүмкін.",
        "Логистикалық тәуекел: қашықтық, жүктің салмағы және көлік түрі маржаны өзгертеді.",
        "Құқықтық тәуекел: тәжірибе, сертификат және техникалық талаптар кәсіпорынға сәйкес келмеуі мүмкін.",
        "Мерзім тәуекелі: кешігу айыппұлға және шарттық беделдің төмендеуіне әкеледі.",
    ])
    add_body(doc, "Қолданыстағы заң бойынша төлем мерзімі, әдетте, міндеттемелер орындалған күннен бастап отыз күннен аспауға тиіс. Орындауды қамтамасыз ету мөлшері қолданылатын рәсімге қарай өзгеруі мүмкін [1]. Сондықтан модельдегі 1 және 3 пайыздық мәндер әмбебап заң талабы ретінде емес, құжаттан нақты мөлшер табылмаған кезде қолданылатын бастапқы параметр ретінде қарастырылады.")
    add_heading(doc, "1.2 Нарықтағы цифрлық шешімдер", level=2)
    add_body(doc, "Салыстыру тек компаниялардың ресми сайттарында 2026 жылғы 14 қыркүйекте ашық көрсетілген функцияларға негізделді. «Жоқ» деген абсолютті баға орнына «ашық сипаттамада көрсетілмеген» деген белгі қолданылды. Бұл тәсіл жасырын немесе кейін қосылған мүмкіндіктер туралы негізсіз қорытынды жасамауға мүмкіндік береді.")
    add_table(
        doc,
        ["Мүмкіндік", "Smart Tender", "Agashka", "Qazaq Tenders"],
        [
            ["Тендерді автоматты бақылау", "Көрсетілген", "Көрсетілген", "Іске асырылған"],
            ["Техникалық ерекшелікті AI арқылы талдау", "Көрсетілген", "Көрсетілген", "Іске асырылған"],
            ["Жеткізушілер ұсыныстары", "Көрсетілген", "Ашық бетте анық емес", "Live API немесе DEMO"],
            ["Кәсіпорын капиталы бойынша күнделікті Cash Flow", "Ашық бетте көрсетілмеген", "Ашық бетте көрсетілмеген", "Негізгі функция"],
            ["Көлік түрлерін салыстыру", "Ашық бетте анық емес", "Ашық бетте көрсетілмеген", "Фура Газель теміржол авиа"],
            ["Параметрлік сценарий", "Ашық бетте анық емес", "Ашық бетте анық емес", "Бір әрекетпен қайта есептеу"],
            ["Нәтиженің формуламен түсіндірілуі", "Толық формула жарияланбаған", "Толық формула жарияланбаған", "Салмақтар мен компоненттер ашық"],
        ],
        [5.4, 3.4, 3.4, 4.4],
        font_size=8.8,
    )
    add_caption(doc, "1 кесте  Нарықтағы сервистердің ашық жарияланған мүмкіндіктері")
    add_body(doc, "Qazaq Tenders-тің айырмашылығы іздеу функциясында емес. Негізгі айырмашылық кәсіпорынның жеке капиталы мен төлем күнтізбесін күндер бойынша модельдеуінде және әр нәтиженің қай деректен алынғанын көрсетуінде.")


def add_methodology(doc: Document) -> None:
    add_heading(doc, "2 ЗЕРТТЕУ ӘДІСТЕМЕСІ", new_page=True)
    add_heading(doc, "2.1 Зерттеу дизайны", level=2)
    add_body(doc, "Зерттеу қолданбалы жобалау және бағдарламалық эксперимент тәсілімен орындалды. Алдымен мемлекеттік сатып алуға қатысу туралы шешімге әсер ететін айнымалылар анықталды. Одан кейін құжаттан алынатын фактілер үшін бірыңғай JSON құрылымы жасалды. Үшінші кезеңде экономикалық модель бағдарламаланды. Соңғы кезеңде сценарийлік есептер мен автоматтандырылған тесттер орындалды.")
    add_table(doc, ["Кезең", "Әдіс", "Нәтиже"], [
        ["Мәселені анықтау", "Заңды және ресми статистиканы талдау", "Тәуекелдер мен айнымалылар тізімі"],
        ["Деректерді құрылымдау", "NLP және қатаң JSON сызбасы", "TenderSpec деректер моделі"],
        ["Экономикалық модель", "Формализация және күнделікті симуляция", "Profit Margin Cash Flow TOS"],
        ["Прототип", "TypeScript Python және aiogram", "Веб интерфейс пен Telegram бот"],
        ["Тексеру", "Модульдік интеграциялық және паритет тесттері", "450 автоматтандырылған тест"],
    ], [3.0, 6.0, 7.6], font_size=9.2)
    add_caption(doc, "2 кесте  Зерттеу кезеңдері")
    add_heading(doc, "2.2 Айнымалылар және дереккөздер", level=2)
    add_body(doc, "Модельдегі айнымалылар үш топқа бөлінді. Бірінші топқа құжаттан тікелей алынатын шарттар жатады: келісімшарт сомасы, жеткізу мерзімі, төлемді кейінге қалдыру, аванс, айыппұл, талап етілетін тәжірибе және сертификаттар. Екінші топ кәсіпорын бейінінен алынады: айналым капиталы, базалық қала, салық режимі, операциялық шығындар, тәжірибе және сертификаттар. Үшінші топ нарықтық бағалаудан тұрады: сатып алу құны, жүк салмағы және логистика тарифі.")
    add_body(doc, "Әр негізгі мәнге provenance, яғни шығу тегі тіркеледі. Құжатта бар сан document, нақты серіктес API-ынан алынған баға supplier API, санаттық бағалау category default, ал пайдаланушы түзеткен мән user verified белгісін алады. Confidence Level осы кірістердің толықтығын көрсетеді. Ол пайданың статистикалық дәлдігі ретінде түсіндірілмейді.")
    add_heading(doc, "2.3 Бағалау өлшемдері", level=2)
    add_bullets(doc, [
        "Функционалдық дұрыстық: формулалардың күтілетін нәтиже беруі.",
        "Қайталанғыштық: бірдей кіріс бірдей нәтиже беруі.",
        "Платформалар паритеті: веб және Telegram қозғалтқыштарының сәйкестігі.",
        "Дерек сапасы: нақты факт пен Smart Default мәнінің ажыратылуы.",
        "Түсіндірілушілік: TOS компоненттері мен тәуекел себептерінің көрсетілуі.",
    ])


def add_architecture(doc: Document) -> None:
    add_heading(doc, "3 QAZAQ TENDERS ЖҮЙЕСІНІҢ АРХИТЕКТУРАСЫ", new_page=True)
    add_picture(doc, WORK / "architecture.png", 16.2, "1 сурет  Qazaq Tenders жүйесіндегі деректер қозғалысы")
    add_heading(doc, "3.1 Деректерді қабылдау", level=2)
    add_body(doc, "Жүйе PDF және DOCX файлдарын, веб сілтемесін және мемлекеттік сатып алу API деректерін қабылдайды. Ресми Goszakup OWS интерфейсі үшін оператор берген токен қажет. Токен болмаса жүйе демо лоттарды бөлек таңбамен көрсетеді. Бұл шектеу рефераттағы «барлық лот автоматты түрде базаға түседі» деген абсолютті тұжырымның орнына нақты жағдайды сипаттайды.")
    add_heading(doc, "3.2 AI қабаты", level=2)
    add_body(doc, "AI қабаты құжат беттерін белгілерімен бірге өңдеп, қатаң JSON сызбасына сәйкес фактілерді шығарады. Модельге белгісіз мәнді факт ретінде ойдан құрастыруға тыйым салынады. Құжатта жоқ сатып алу құны мен жүк салмағы ғана Smart Default ретінде бағалануы мүмкін. Жеткізушіге қойылатын арнайы талаптар базалық құжаттардан бөлек шығарылады және бастапқы бет нөмірімен сақталады.")
    add_bullets(doc, [
        "жеткізу және төлем мерзімдері" ,
        "аванс кепілдік және айыппұл шарттары",
        "лицензия сертификат тәжірибе персонал және жабдық талаптары",
        "тауардың техникалық сипаттамалары мен кепілдік мерзімі",
        "бәсекені шектеуі ықтимал талаптар және олардың құжаттағы беті",
    ])
    add_heading(doc, "3.3 Экономикалық қабат", level=2)
    add_body(doc, "Экономикалық қабатта тілдік модель қолданылмайды. Барлық ақшалай нәтиже формулалар арқылы есептеледі. Осы бөлу тілдік модельдің арифметикалық қатесін TOS пен Cash Flow нәтижесіне тікелей өткізбеуге мүмкіндік береді. Пайдаланушы көлік түрін, авансты немесе жеткізуші бағасын өзгерткен сайын сол алгоритм қайта орындалады.")
    add_heading(doc, "3.4 Нәтижені ұсыну", level=2)
    add_body(doc, "Веб интерфейс пен Telegram-бот бір деректер келісімін және бірдей формулаларды қолданады. Карточкада TOS, болжамды таза пайда, маржа, кассалық алшақтық, логистика, дерек сенімділігі және негізгі тәуекел көрсетіледі. Толық бөлімде күнделікті ақша ағыны мен формула компоненттері ашылады.")


def add_model(doc: Document) -> None:
    add_heading(doc, "4 ЭКОНОМИКАЛЫҚ МАТЕМАТИКАЛЫҚ МОДЕЛЬ", new_page=True)
    add_heading(doc, "4.1 Пайда және маржа", level=2)
    add_body(doc, "Келісімшарт бойынша болжамды таза пайда кірістен барлық есептелетін шығындарды шегеру арқылы табылады. Модель салықты, банктік кепілдік комиссиясын, кассалық алшақтықты жабуға арналған кредит пайызын және кешігу айыбын бөлек есептейді.")
    add_equation(doc, [_mr("Π = S − ("), _sub("C", "purchase"), _mr(" + "), _sub("C", "logistics"), _mr(" + "), _sub("C", "tax"), _mr(" + "), _sub("C", "bank"), _mr(" + "), _sub("C", "oper"), _mr(" + "), _sub("C", "penalty"), _mr(")")], "1 формула  Болжамды таза пайда")
    add_equation(doc, [_sub("M", "rel"), _mr(" = "), _frac([_mr("Π")], [_mr("S")]), _mr(" · 100%     "), _sub("M", "score"), _mr(" = min(100, max(0, "), _frac([_sub("M", "rel")], [_mr("20%")]), _mr(" · 100))")], "2 формула  Маржа және нормаланған маржа балы")
    add_body(doc, "Мұнда S келісімшарт сомасын, ал C индекстері сатып алу, логистика, салық, банк, операциялық және айыппұл шығындарын білдіреді. TOS ішінде Mrel тікелей қолданылмайды. 20 пайыз және одан жоғары маржа 100 балға тең деп алынып, маржа көрсеткіші 0 мен 100 аралығына нормаланады. Бұл шек модельдік параметр болып саналады және кейін эмпирикалық дерекпен калибрленуге тиіс.")
    add_heading(doc, "4.2 Күнделікті ақша ағыны", level=2)
    add_body(doc, "Пайданың оң болуы өтімділіктің жеткілікті екенін білдірмейді. Сондықтан модель шарт аяқталып, соңғы төлем түскенге дейін әр күндегі қалдықты есептейді.")
    add_equation(doc, [_sub("CF", "t"), _mr(" = "), _sub("CF", "0"), _mr(" + Σ Inflowᵢ − Σ Outflowᵢ     Gap = 1{min "), _sub("CF", "t"), _mr(" < 0}")], "3 формула  Күнделікті ақша ағыны және кассалық алшақтық шарты")
    add_equation(doc, [_sub("CF", "risk"), _mr(" = clamp((1 − "), _frac([_sub("CF", "min")], [_sub("CF", "0")]), _mr(") · 45), егер "), _sub("CF", "min"), _mr(" ≥ 0")], "4 формула  Оң қалдық кезіндегі Cash Flow тәуекелі")
    add_equation(doc, [_sub("CF", "risk"), _mr(" = clamp(50 + 0,35D + 0,15T), егер "), _sub("CF", "min"), _mr(" < 0")], "5 формула  Кассалық алшақтық кезіндегі тәуекел")
    add_body(doc, "D көрсеткіші кассалық тапшылық тереңдігінің бастапқы капиталға қатынасын, ал T тапшылық болған күндердің бүкіл кезеңге қатынасын білдіреді. clamp функциясы нәтижені 0 мен 100 арасында шектейді. Егер қалдық теріс болмаса, тәуекел ақша жастығының қаншалық азайғанына қарай есептеледі.")
    add_heading(doc, "4.3 Логистикалық және құқықтық тәуекел", level=2)
    add_equation(doc, [_sub("L", "score"), _mr(" = max(0, 100 − "), _frac([_mr("Dist")], [_sub("Dist", "max")]), _mr(" · 100)     Pen = min(S · "), _sub("K", "delay"), _mr(" · "), _sub("d", "late"), _mr(", 0,10S)")], "6 формула  Логистика балы және кешігу айыбы")
    add_body(doc, "Dist кәсіпорын базасынан жеткізу орнына дейінгі есептік қашықтық, Distmax кәсіпорын профиліндегі тиімді радиус. Айыппұл формуласы шарттағы тәуліктік мөлшерлемені қолданады және модельдік әдепкі шекті 10 пайызбен шектейді. Нақты сатып алуда құжаттағы мөлшерлеме мен шек басымдыққа ие.")
    add_body(doc, "Құқықтық тәуекел тәжірибе жеткіліксіздігі, сертификаттардың болмауы, техникалық ерекшеліктегі жоғары тәуекелді талаптар және кешігу сценарийі бойынша жинақталады. Бұл бал заңдық қорытынды емес. Ол кәсіпкерге тексеруді қажет ететін тармақтарды ерте көрсетуге арналған.")
    add_heading(doc, "4.4 Tender Opportunity Score", level=2)
    add_equation(doc, [_mr("TOS = 0,35"), _sub("M", "score"), _mr(" + 0,30(100 − "), _sub("CF", "risk"), _mr(") + 0,15"), _sub("L", "score"), _mr(" + 0,20(100 − "), _sub("R", "legal"), _mr(")")], "7 формула  Tender Opportunity Score индексі")
    add_table(doc, ["Компонент", "Салмақ", "Неліктен енгізілді"], [
        ["Маржа балы", "0,35", "Экономикалық нәтиженің негізгі көрсеткіші"],
        ["Cash Flow қауіпсіздігі", "0,30", "ШОБ үшін өтімділік тәуекелі жоғары"],
        ["Логистика", "0,15", "Қашықтық пен жеткізу қолжетімділігін есепке алу"],
        ["Құқықтық қауіпсіздік", "0,20", "Талаптарға сәйкессіздік пен санкция тәуекелі"],
    ], [4.0, 2.2, 10.4], font_size=9.5)
    add_caption(doc, "3 кесте  TOS индексінің салмақтары")
    add_body(doc, "Салмақтар сараптамалық бастапқы нұсқа ретінде қабылданды. Олар 1,00-ге тең, бірақ статистикалық оңтайландырудан өтпеген. Сондықтан TOS қатысу туралы соңғы шешімді алмастырмайды. Индекс әртүрлі сценарийлерді бірдей шкалада салыстыруға арналған шешім қолдау құралы.")


def add_logistics(doc: Document) -> None:
    add_heading(doc, "5 ЛОГИСТИКА ЖӘНЕ ЖЕТКІЗУШІЛЕР МОДУЛІ", new_page=True)
    add_heading(doc, "5.1 Қашықтық пен көлік режимдері", level=2)
    add_body(doc, "Жеткізу қаласы техникалық ерекшеліктен, ал кәсіпорынның базалық қаласы пайдаланушы профилінен алынады. Қалалар сәйкес келсе, қала ішіндегі тұрақты тариф қолданылады. Қалалар әртүрлі болса, Haversine формуласы бойынша түзу қашықтық есептеліп, автомобиль бағыты үшін 1,25, теміржол үшін 1,10 коэффициенті қолданылады. Бұл коэффициенттер маршруттық бағалау болып табылады және навигациялық сервис бағытын алмастырмайды.")
    add_table(doc, ["Көлік", "Базалық тариф", "Сыйымдылық", "Модельдік мерзім"], [
        ["Автофура", "450 ₸/км және 50% кері жол, кемі 60 000 ₸", "20 т", "650 км тәулігіне"],
        ["Газель", "180 ₸/км, кемі 30 000 ₸", "3 т", "550 км тәулігіне"],
        ["Теміржол", "12 000 ₸/т әр 1000 км және терминал", "20 т контейнер", "450 км тәулігіне және 2 күн"],
        ["Авиа", "850 ₸/кг, кемі 60 000 ₸", "Салмақ бойынша", "1 күн"],
        ["Қала ішінде", "25 000 ₸ немесе ауыр рейске 60 000 ₸", "3 т бір рейс", "1 күн"],
    ], [3.0, 7.4, 2.8, 3.4], font_size=8.9)
    add_caption(doc, "4 кесте  Логистикалық тарифтердің бастапқы параметрлері")
    add_body(doc, "Бұл тарифифтер тасымалдаушының коммерциялық офертасы емес. Автомобиль тасымалы үшін ATI.SU API қолжетімді болғанда live нарықтық ставка қолданылады және интерфейсте дереккөз бөлек белгіленеді. API қолжетімсіз болса, жүйе есептік тарифке қауіпсіз ауысады. Теміржол мен авиа тарифтері әзірге тек модельдік бағдар болып қалады.")
    add_heading(doc, "5.2 Жеткізушілер қозғалтқышы", level=2)
    add_body(doc, "Пайдаланушы лот ішінен жеткізушілер тізімін аша алады. Серіктес каталог қосылған жағдайда ұсыныс бағасы, бірлік бағасы, қала, байланыс, жүк салмағы, СТ KZ сертификаты және жаңарту уақыты көрсетіледі. Нақты жауап verified немесе smart AI мәртебесімен белгіленеді.")
    add_body(doc, "Серіктес каталог бапталмаған немесе 503 қатесін қайтарған жағдайда жүйе бес өңір бойынша анық DEMO белгісі бар мысалдарды көрсетеді. Демо атаулар мен байланыстар нақты коммерциялық ұсыныс ретінде ұсынылмайды, Verified мәртебесін алмайды және Confidence Level көрсеткішін көтермейді. Бұл шешім прототипті көрсетуге мүмкіндік береді әрі жалған нарықтық деректі нақты дерек ретінде көрсетуден қорғайды.")


def add_case(doc: Document, results: dict, context: dict[str, object]) -> None:
    spec = context["spec"]
    add_heading(doc, "6 ЕСЕПТІК КЕЙС ЖӘНЕ СЦЕНАРИЙЛІК ТАЛДАУ", new_page=True)
    add_heading(doc, "6.1 Кіріс деректері", level=2)
    add_body(doc, "Сценарийлік тексеру үшін бағдарламаның демо деректеріндегі облыстық мектептерге компьютерлік жабдық жеткізу лоты қолданылды. Кейс нақты конкурс нәтижесін болжау үшін емес, модельдің айнымалыларға реакциясын көрсету үшін таңдалды.")
    add_table(doc, ["Көрсеткіш", "Мән"], [
        ["Келісімшарт сомасы", "84 000 000 ₸"],
        ["Сатып алу құны", "62 000 000 ₸"],
        ["Жүк", "14 т"],
        ["Бағыт", "Астана Талдықорған"],
        ["Жеткізу мерзімі", "30 күн"],
        ["Төлемді кейінге қалдыру", "30 күн"],
        ["Бастапқы айналым капиталы", "52 000 000 ₸"],
        ["Салық режимі", "Жалпы белгіленген режим"],
    ], [7.0, 9.6], font_size=9.5)
    add_caption(doc, "5 кесте  Үлгілік кейстің бастапқы параметрлері")
    add_heading(doc, "6.2 Нәтижелер", level=2)
    rows = []
    for name in ("Автоматты таңдау", "Автофура", "Теміржол", "Авиа", "30% аванс"):
        r = results[name]
        rows.append([
            name,
            f"{r.tos:.1f}",
            f"{r.net_profit / 1e6:.2f}",
            f"{r.margin_pct:.1f}%",
            f"{r.costs.logistics / 1e6:.2f}",
            "Жоқ" if not r.cash_flow_gap else f"{r.max_deficit / 1e6:.2f}",
        ])
    add_table(doc, ["Сценарий", "TOS", "Пайда млн ₸", "Маржа", "Логистика млн ₸", "Алшақтық млн ₸"], rows, [4.0, 1.7, 2.9, 2.2, 3.0, 3.2], font_size=8.7)
    add_caption(doc, "6 кесте  Сценарийлер бойынша есеп нәтижесі")
    add_picture(doc, WORK / "logistics_case.png", 14.7, "2 сурет  Көлік түріне қарай логистика құны")
    add_body(doc, "Авиа жеткізу бір күнге дейін қысқарғанымен, 14 тонна жүк үшін логистика құны 11,90 млн теңгеге жетіп, TOS 41,5 балға дейін төмендеді. Автофураның құны 0,72 млн теңге, ал теміржолдың есептік құны 0,28 млн теңге болды. Бұл нәтиже көлік таңдауын тек жылдамдықпен емес, жүк салмағы және маржамен бірге қарау керектігін көрсетеді.")
    add_picture(doc, WORK / "cashflow_case.png", 15.3, "3 сурет  Аванстың күнделікті ақша ағынына әсері")
    add_body(doc, "Автоматты сценарийде болжамды пайда 14,55 млн теңге және маржа 17,3 пайыз болғанымен, кассалық тапшылық 13,19 млн теңгеге жетті. 30 пайыз аванс енгізілгенде алшақтық жойылып, TOS 73,4 балға өсті. Демек, бұл кейсте қатысу шешімін өзгерткен негізгі фактор пайда көлемі емес, төлем күнтізбесі болды.")
    add_heading(doc, "6.3 Сценарийді түсіндіру", level=2)
    add_body(doc, "Кейс модельдің себеп салдарын дұрыс бағытта көрсететінін дәлелдейді: қымбат тасымал пайданы азайтады, ұзақ жол кешігу тәуекелін өсіреді, ал аванс ерте кезеңдегі қаржыландыру қажеттілігін төмендетеді. Алайда бір демо кейс модельдің нарықтағы болжау дәлдігін анықтамайды. Ол тек алгоритмнің экономикалық логикасын көрсетеді.")


def add_validation(doc: Document) -> None:
    add_heading(doc, "7 БАҒДАРЛАМАЛЫҚ ВАЛИДАЦИЯ", new_page=True)
    add_heading(doc, "7.1 Автоматтандырылған тексеру", level=2)
    add_body(doc, "2026 жылғы 14 қыркүйекте жобаның Python тесттер жиынтығы орындалды. Желіден оқшауланған ортада 446 тест сәтті өтті. ATI.SU сервисіне тікелей қосылатын төрт интеграциялық тест желі рұқсатымен бөлек орындалып, төртеуі де сәтті аяқталды. Осылайша тексерілген тесттердің жалпы саны 450 болды.")
    add_table(doc, ["Тексеру түрі", "Қамту", "Нәтиже"], [
        ["Формула тесттері", "Пайда салық кепілдік айыппұл және Cash Flow", "Сәтті"],
        ["Шекаралық жағдайлар", "Нөлдік қашықтық үлкен жүк кешігу және теріс маржа", "Сәтті"],
        ["TypeScript Python паритеті", "210 бірдей лот профиль сценарий комбинациясы", "Сәйкес"],
        ["Интерфейс тесттері", "Telegram карточкалары батырмалар және тілдер", "Сәтті"],
        ["API fallback", "Goszakup жеткізушілер және логистика қатесі", "Сәтті"],
        ["ATI.SU live API", "Екі бағыт және екі автомобиль режимі", "4 тест сәтті"],
    ], [4.1, 9.0, 3.5], font_size=9.1)
    add_caption(doc, "7 кесте  Бағдарламалық валидация нәтижелері")
    add_body(doc, "Паритет тестінің мәні ерекше. Бірдей 210 кіріс комбинациясы веб қозғалтқышында алдын ала есептеліп, Python нұсқасының нәтижесімен салыстырылды. TOS, пайда, маржа, кассалық алшақтық, күндер, тәуекел компоненттері және шығындар рұқсат етілген дөңгелектеу шегінде сәйкес келді.")
    add_heading(doc, "7.2 Валидация нені дәлелдейді", level=2)
    add_body(doc, "Автоматтандырылған тесттер бағдарламаның сипатталған формулаларды тұрақты орындайтынын және екі платформада бірдей нәтиже беретінін дәлелдейді. Олар Smart Default бағасының нақты нарықтық бағамен қаншалық сәйкес келетінін немесе NLP модулінің барлық техникалық талапты табу дәлдігін өлшемейді.")
    add_heading(doc, "7.3 Эмпирикалық тексеру жоспары", level=2)
    add_body(doc, "Экономикалық және NLP дәлдігін бағалау үшін келесі кезеңде кемінде 100 аяқталған тендерден тұратын деректер жиынтығы қажет. Құжаттарды екі тәуелсіз сарапшы белгілеп, келіспеген тармақтарды үшінші сарапшы тексеруі тиіс.")
    add_table(doc, ["Мақсат", "Метрика", "Салыстыру дерегі"], [
        ["Талаптарды шығару", "Precision Recall F1", "Сарапшы белгілеген тармақтар"],
        ["Сатып алу құны", "MAPE және медианалық абсолют қате", "Нақты шот немесе коммерциялық ұсыныс"],
        ["Логистика құны", "MAPE", "Тасымалдаушының нақты шоты"],
        ["Cash Flow", "Күн мен тапшылық сомасының қатесі", "Банк көшірмесі және төлем күнтізбесі"],
        ["TOS шешімі", "Дәлдік және ROC AUC", "Шарттың нақты пайдасы мен орындалуы"],
        ["Талдау уақыты", "Медианалық минут және үнем пайызы", "Қолмен және жүйемен орындалған тапсырма"],
    ], [4.1, 5.0, 7.5], font_size=9.0)
    add_caption(doc, "8 кесте  Келесі эмпирикалық валидацияның метрикалары")
    add_body(doc, "Осы эксперимент орындалмайынша жұмыста 95 пайыз дәлдік немесе талдау уақытын 90 пайыз қысқарту туралы қорытынды жасалмайды. Мұндай көрсеткіштер тек өлшенген іріктеме, бастапқы деңгей және сенімділік аралығы берілгеннен кейін жариялануы тиіс.")


def add_limitations(doc: Document) -> None:
    add_heading(doc, "8 ШЕКТЕУЛЕР ЖӘНЕ ДАМЫТУ БАҒЫТТАРЫ", new_page=True)
    add_heading(doc, "8.1 Қазіргі шектеулер", level=2)
    add_numbered(doc, [
        "Goszakup хабарландыру API-ы техникалық ерекшеліктің барлық шартын бермейді. Толық талдау үшін PDF немесе DOCX қажет.",
        "Smart Default мәндері санаттық коэффициенттерге сүйенеді және нақты жеткізушінің коммерциялық ұсынысын алмастырмайды.",
        "Теміржол мен авиа тарифтері есептік бағдар болып саналады. Live тариф әзірге автомобиль тасымалына ғана қосылған.",
        "Haversine коэффициенті нақты жол жабылуын, маусымдық маршрутты және терминалға дейінгі соңғы шақырымды ескермейді.",
        "TOS салмақтары сараптамалық түрде таңдалған және тарихи жеңіс немесе пайда деректерімен калибрленбеген.",
        "Құқықтық тәуекел балы заңгер қорытындысын алмастырмайды.",
        "Демо жеткізушілер нақты компания немесе баға ретінде пайдаланылмауға тиіс.",
    ])
    add_heading(doc, "8.2 Дамыту бағыты", level=2)
    add_body(doc, "Бірінші кезектегі міндет нақты аяқталған шарттар бойынша валидациялық деректер жиынтығын қалыптастыру. Екінші міндет жеткізушілердің келісілген каталогтарын және теміржол мен авиа тасымалдаушыларының коммерциялық API интерфейстерін қосу. Үшінші міндет TOS салмақтарын тарихи дерек бойынша калибрлеу және әр нәтиже үшін сенімділік аралығын көрсету.")
    add_body(doc, "Өнімдік даму тұрғысынан кәсіпорынның банк көшірмесін пайдаланушы келісімімен импорттау, бірнеше шарттың бір мезгілдегі ақша ағынын модельдеу және нақты орындалған шарттан кейін жоспар мен факт айырмасын есептеу маңызды. Бұл өзгерістер модельді бір тендерді бағалау құралынан компанияның тендерлік портфелін басқару жүйесіне айналдырады.")


def add_conclusion(doc: Document) -> None:
    add_heading(doc, "ҚОРЫТЫНДЫ", new_page=True)
    add_body(doc, "Зерттеу барысында мемлекеттік сатып алуға қатысудың экономикалық тиімділігін келісімшарт сомасы арқылы ғана бағалау жеткіліксіз екені көрсетілді. Пайдасы оң лоттың өзі төлем күнтізбесі кәсіпорын капиталымен сәйкес келмесе, кассалық алшақтық туғызады.")
    add_body(doc, "Осы мәселені шешу үшін Qazaq Tenders прототипі әзірленді. Жүйе құжаттан талаптарды шығаратын AI қабатын ақшалай есеп жүргізетін детерминирленген экономикалық қабаттан бөледі. Модель пайда, маржа, күнделікті Cash Flow, логистика және құқықтық тәуекел компоненттерін түсіндірілетін TOS индексіне біріктіреді.")
    add_body(doc, "Үлгілік есеп 30 пайыз аванс кассалық тапшылықты жойып, TOS көрсеткішін 62,7-ден 73,4 балға көтергенін көрсетті. Авиа тасымалы мерзімді қысқартқанымен, ауыр жүк үшін маржаны айтарлықтай азайтты. Бұл нәтижелер сценарийлік модельдің кәсіпкерге шарт талаптарының салдарын өтінімге дейін салыстыруға көмектесетінін көрсетеді.")
    add_body(doc, "Бағдарламалық тексеруде 450 тест сәтті орындалды, оның ішінде TypeScript пен Python қозғалтқыштарының 210 сценарий бойынша паритеті расталды. Бұл нәтиже есептеудің техникалық тұрақтылығын дәлелдейді. Экономикалық болжамның нарықтық дәлдігі мен NLP сапасы бөлек эмпирикалық зерттеуді қажет етеді.")
    add_body(doc, "Зерттеу гипотезасы бағдарламалық және сценарийлік деңгейде қолдау тапты: айналым капиталы мен төлем күнтізбесін модельге енгізу пайдалы көрінетін, бірақ қаржыландыру тұрғысынан қауіпті тендерді анықтауға мүмкіндік береді. Келесі ғылыми кезең нақты аяқталған тендерлер бойынша модельді калибрлеу мен тәуелсіз сараптамалық бағалаудан тұрады.")


def add_references(doc: Document) -> None:
    add_heading(doc, "ПАЙДАЛАНЫЛҒАН ӘДЕБИЕТТЕР", new_page=True)
    refs = [
        "Қазақстан Республикасының 2024 жылғы 1 шілдедегі № 106 VIII «Мемлекеттік сатып алу туралы» Заңы. Әділет ақпараттық құқықтық жүйесі. https://adilet.zan.kz/rus/docs/Z2400000106",
        "Қазақстан Республикасы Қаржы министрлігі. 2025 жылы мемлекеттік сатып алу көлемі 8,1 трлн теңгеге дейін өсті. https://www.gov.kz/memleket/entities/minfin/press/news/details/1151855",
        "Қазақстан Республикасы Ұлттық статистика бюросы. Қазақстан Республикасындағы шағын және орта кәсіпкерліктің мониторингі. 2026 жылғы 1 қаңтар. https://stat.gov.kz/ru/industries/businessstatistics/stat-org/publications/306675/",
        "Қазақстан Республикасының мемлекеттік сатып алу веб порталы. https://goszakup.gov.kz",
        "Мемлекеттік сатып алудың ашық веб сервисі OWS v3. https://ows.goszakup.gov.kz",
        "OECD. Public Procurement in Kazakhstan Reforming for Efficiency. OECD Publishing, 2019. https://doi.org/10.1787/c11183ae-en",
        "World Bank. Procurement Regulations for IPF Borrowers. Washington DC, 2025. https://www.worldbank.org/en/projects-operations/products-and-services/brief/procurement-new-framework",
        "NIST. Artificial Intelligence Risk Management Framework AI RMF 1.0. 2023. https://doi.org/10.6028/NIST.AI.100-1",
        "OpenAI. Structured Outputs documentation. https://platform.openai.com/docs/guides/structured-outputs",
        "Sinnott R. W. Virtues of the Haversine. Sky and Telescope. 1984. Vol. 68. No. 2. P. 159.",
        "Brealey R. Myers S. Allen F. Principles of Corporate Finance. 14th ed. McGraw Hill, 2022.",
        "NumPy Developers. NumPy documentation. https://numpy.org/doc/",
        "Python Software Foundation. Python documentation. https://docs.python.org/3/",
        "Smart Tender. Платформа закупок Республики Казахстан. https://www.smarttender.kz/",
        "Agashka. Тендерная аналитика для Казахстана. https://agashkabot.kz/ru/",
        "Qazaq Tenders жобасының бастапқы коды және автоматтандырылған тесттері. Авторлық бағдарламалық материал, 2026.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        set_font(p.add_run(f"{i}. {ref}"), size=10.5)
    add_body(doc, "Интернет дереккөздеріне жүгінген күн: 2026 жылғы 14 қыркүйек.", align=WD_ALIGN_PARAGRAPH.LEFT)


def add_appendix(doc: Document) -> None:
    add_heading(doc, "ҚОСЫМША", new_page=True)
    add_heading(doc, "А қосымшасы  Деректер келісімінің негізгі өрістері", level=2)
    add_table(doc, ["Өріс", "Мағынасы", "Негізгі дереккөз"], [
        ["contractAmount", "Келісімшарт сомасы", "Хабарландыру немесе құжат"],
        ["purchaseCost", "Сатып алу немесе өндіріс құны", "Жеткізуші API немесе Smart Default"],
        ["cargoTonnes", "Жүктің жалпы салмағы", "Құжат немесе санаттық бағалау"],
        ["deliveryDays", "Жеткізу мерзімі", "Техникалық ерекшелік"],
        ["paymentDelayDays", "Төлемді кейінге қалдыру", "Шарт жобасы"],
        ["advancePercentage", "Аванс үлесі", "Шарт жобасы немесе пайдаланушы"],
        ["hiddenRequirements", "Тәуекелді немесе шектеуші талаптар", "Бет нөмірі бар NLP нәтижесі"],
        ["fieldSources", "Әр кірістің шығу тегі мен сенімділігі", "Жүйелік provenance"],
    ], [4.0, 7.4, 5.2], font_size=9.1)
    add_caption(doc, "А.1 кесте  TenderSpec құрылымының негізгі өрістері")
    add_heading(doc, "Ә қосымшасы  Нәтижені түсіндіру ережесі", level=2)
    add_bullets(doc, [
        "TOS 70 және одан жоғары болса да, таза пайда теріс лот ұсынылмайды.",
        "Кассалық алшақтық бар лот автоматты түрде қауіпсіз деп белгіленбейді.",
        "Smart Default мәндері интерфейсте бағалау ретінде көрсетіледі.",
        "Verified мәртебесі тек пайдаланушы растаған немесе сенімді API берген дерекке қолданылады.",
        "Демо лот пен демо жеткізуші нақты нарықтық дерек ретінде көрсетілмейді.",
    ])
    add_heading(doc, "Б қосымшасы  Жобаны қорғауда көрсетуге ұсынылатын сценарий", level=2)
    add_numbered(doc, [
        "Кәсіпорын профилінде Астана қаласын және 52 млн теңге айналым капиталын көрсету.",
        "84 млн теңгелік компьютер жабдығы лотын ашу.",
        "TOS пен 13,19 млн теңгелік кассалық алшақтықты түсіндіру.",
        "Көлік режимдерін ауыстырып, пайда мен логистика өзгерісін көрсету.",
        "30 пайыз аванс батырмасын басып, кассалық алшақтықтың жойылғанын көрсету.",
        "Құжат талаптарының бастапқы бет нөмірін және Smart Default белгісін көрсету.",
    ])


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    results, context = make_visuals()
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_front_matter(doc)
    add_introduction(doc)
    add_market(doc)
    add_methodology(doc)
    add_architecture(doc)
    add_model(doc)
    add_logistics(doc)
    add_case(doc, results, context)
    add_validation(doc)
    add_limitations(doc)
    add_conclusion(doc)
    add_references(doc)
    add_appendix(doc)

    props = doc.core_properties
    props.title = "Қазақстандағы ШОБ үшін мемлекеттік сатып алудың экономикалық тиімділігін бағалау"
    props.subject = "Qazaq Tenders ғылыми жобасы"
    props.author = "Байбосынов Нұрали Айбарұлы"
    props.keywords = "мемлекеттік сатып алу, ШОБ, TOS, Cash Flow, NLP"
    props.comments = "2026 жылғы конкурсқа арналған ғылыми жоба"
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
